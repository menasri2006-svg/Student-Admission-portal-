import os
import sys

# Complete Word Report Generator for Student Admission Portal
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

out_dir = r"E:\Webproject\submission"
os.makedirs(out_dir, exist_ok=True)
docx_path = os.path.join(out_dir, "Student_Admission_Portal_Final_Report.docx")
md_path = os.path.join(out_dir, "Student_Admission_Portal_Final_Report.md")

screenshots_base = r"E:\Webproject\submission\screenshots"
co3_dir = os.path.join(screenshots_base, "CO3_Registration")
co4_dir = os.path.join(screenshots_base, "CO4_PHP_CRUD")
co5_dir = os.path.join(screenshots_base, "CO5_Servlet_JSP")

# Styles and Helpers
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Student Admission Portal  |  Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(100, 116, 139)
    
    fld1 = OxmlElement('w:fldSimple')
    fld1.set(qn('w:instr'), 'PAGE')
    p._p.append(fld1)
    
    run2 = p.add_run(" of ")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(100, 116, 139)
    
    fld2 = OxmlElement('w:fldSimple')
    fld2.set(qn('w:instr'), 'NUMPAGES')
    p._p.append(fld2)

def h1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p

def h2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 58, 138)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p

def h3(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    return p

def p_body(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_main = p.add_run(text)
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(12)
    r_main.font.italic = italic
    r_main.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p

def p_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_main = p.add_run(text)
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(12)
    r_main.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    return p

def code_box(doc, code_str):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    set_table_borders(tbl, color="CBD5E1", sz="4", val="single")
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_str.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def figure(doc, img_path, caption_text, width_in=5.8):
    if not os.path.exists(img_path):
        print(f"Warning: Figure not found at {img_path}")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture(img_path, width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(caption_text)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)

def table_grid(doc, headers, data_rows, col_widths=None):
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="4", val="single")
    
    # Header row
    for col_idx, h in enumerate(headers):
        cell = tbl.cell(0, col_idx)
        if col_widths and col_idx < len(col_widths):
            cell.width = col_widths[col_idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(h)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data rows
    for row_idx, row_data in enumerate(data_rows):
        fill_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = tbl.cell(row_idx + 1, col_idx)
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
            set_cell_background(cell, fill_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(str(val))
            run.font.name = "Times New Roman"
            run.font.size = Pt(9.5)
            if val == "PASS" or val == "PASSED":
                run.font.bold = True
                run.font.color.rgb = RGBColor(22, 101, 52)
            elif val == "FAIL" or val == "FAILED":
                run.font.bold = True
                run.font.color.rgb = RGBColor(185, 28, 28)
            else:
                run.font.color.rgb = RGBColor(30, 41, 59)
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

print("Core functions loaded.")
