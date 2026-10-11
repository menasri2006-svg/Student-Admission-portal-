import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_footer_page_numbers(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Student Admission Portal  |  Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(120, 120, 120)
    
    fld1 = OxmlElement('w:fldSimple')
    fld1.set(qn('w:instr'), 'PAGE')
    p._p.append(fld1)
    
    run2 = p.add_run(" of ")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(120, 120, 120)
    
    fld2 = OxmlElement('w:fldSimple')
    fld2.set(qn('w:instr'), 'NUMPAGES')
    p._p.append(fld2)

def style_paragraph(p, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color_rgb=(0,0,0), space_before=0, space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    for run in p.runs:
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = RGBColor(*color_rgb)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    style_paragraph(p, font_name="Times New Roman", size_pt=14, bold=True, color_rgb=(15, 23, 42), space_before=14, space_after=6)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    style_paragraph(p, font_name="Times New Roman", size_pt=12, bold=True, color_rgb=(30, 58, 138), space_before=10, space_after=4)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    style_paragraph(p, font_name="Times New Roman", size_pt=11, bold=True, italic=True, color_rgb=(51, 65, 85), space_before=6, space_after=2)
    return p

def add_body_p(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_main = p.add_run(text)
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(12)
    r_main.font.italic = italic
    r_main.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    return p

def add_bullet_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(15, 23, 42)
    r_main = p.add_run(text)
    r_main.font.name = "Times New Roman"
    r_main.font.size = Pt(12)
    r_main.font.color.rgb = RGBColor(30, 41, 59)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    return p

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    # spacing after table
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_image_figure(doc, img_path, fig_title, width_in=6.0):
    if not os.path.exists(img_path):
        print(f"Warning: Image not found at {img_path}")
        return
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(6)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run().add_picture(img_path, width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    r_cap = p_cap.add_run(fig_title)
    r_cap.font.name = "Times New Roman"
    r_cap.font.size = Pt(10)
    r_cap.font.bold = True
    r_cap.font.color.rgb = RGBColor(71, 85, 105)

print("Formatting primitives compiled.")
